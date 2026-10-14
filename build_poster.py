"""Builds the A4 poster (shrinkflation_poster.pptx).

Each chart has a takeaway line above it and a short interpretation below it. The charts come from figs_poster/
(made by make_poster_figs.py), drawn shorter than the notebook's to leave room for the interpretation.

Run: uv run --no-project --with python-pptx python build_poster.py
"""

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

NAVY, BLUE2, ORANGE, INK, INK2, MUTED = "0d366b", "1c5cab", "eb6834", "0b0b0b", "52514e", "6f6e69"
CARD, INTRO, WHITE = "f4f5f7", "eef4fc", "ffffff"
FONT = "Arial"

SW, SH, M = 8.27, 11.69, 0.3
CW = SW - 2 * M
GUT = 0.14
COL = (CW - 2 * GUT) / 3
FIG_W, FIG_H = 2.36, 1.7
FIG_DIR = "figs_poster"
INSIGHT_H = 0.47

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
slide = prs.slides.add_slide(prs.slide_layouts[6])

Run = tuple[str, dict[str, object]]


def box(x: float, y: float, w: float, h: float, fill: str, name: str, shape: MSO_SHAPE = MSO_SHAPE.ROUNDED_RECTANGLE) -> None:
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.04


def text(x: float, y: float, w: float, h: float, paras: list[tuple[list[Run], bool]], name: str, size: float = 6.5,
         color: str = INK, anchor: MSO_ANCHOR = MSO_ANCHOR.TOP, space_after: float = 0, bullet_color: str = ORANGE) -> None:
    """paras: (runs, is_bullet) per paragraph. Run options: b, color, size."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, (runs, bullet) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.05
        p.space_after = Pt(space_after)
        if bullet:   # real bullet with a hanging indent
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.1))))
            pPr.set("indent", str(-Emu(Inches(0.1))))
            clr = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(clr, qn("a:srgbClr")).set("val", bullet_color)
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", FONT)
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
        for t, o in runs:
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(float(o.get("size", size)))
            r.font.bold = bool(o.get("b", False))
            r.font.color.rgb = RGBColor.from_string(str(o.get("color", color)))


def bullets(items: list[str], color: str = INK) -> list[tuple[list[Run], bool]]:
    return [([(t, {"color": color})], True) for t in items]


def bullets_lead(items: list[tuple[str, str]], color: str = INK) -> list[tuple[list[Run], bool]]:
    return [([(lead, {"b": True, "color": color}), (rest, {"color": color})], True) for lead, rest in items]


# ---------- title block: what is explored ----------
LIGHT = "8f8e89"
text(M, 0.18, CW, 0.42, [([("Shrinkflation in Singapore's Private Homes", {"b": True})], False)], "Title", size=23, color=NAVY)
text(M, 0.6, CW, 0.14, [([("DAO2702 Group Assignment  ·  Team Jupyter  ·  Martin, Huaitao, Manan, Jialin, Midhvan, Hoyoung", {})], False)], "Team line", size=6.5, color=MUTED)
text(M, 0.8, CW, 0.24, [([("Did homes shrink to keep the price tag within reach, and is it still happening in 2025?", {})], False)],
     "Question", size=13, color=INK)
# the few terms needed to read the charts, two quiet lines
text(M, 1.1, CW, 0.22, [
    ([("Shrinkflation: less product for the same price tag [6]   ·   PSF: price per sq ft, so price tag = floor area × PSF   ·   "
       "Lease-start year ≈ year the land was sold", {})], False),
    ([("URA unit cap: units ≤ floor area ÷ 70 m² (85 m² from 2019) [5]   ·   MAS limit: loan repayments ≤ 55% of income, tested at 4% [4]", {})], False),
], "Terms line", size=6.2, color=LIGHT)
KY, KH = 1.1, 0.24   # rows start below the terms line

# ---------- panels: one takeaway per chart ----------
ROWS = ["WHAT SHRANK", "IS IT REAL, AND HOW DOES IT WORK?", "WHY IT HAPPENED, AND WHAT CHANGED"]
PANELS = [
    ("01_histogram.png", "New launches are 21% smaller than resale homes"),
    ("02_line_vintage.png", "Homes shrank 43% for land sold after 2008"),
    ("03_stacked_mix.png", "Shoeboxes rose from 1% to 13%, big homes fell from 34% to 8%"),
    ("04_coefficients.png", "Within the same district the gap stays at −34%"),
    ("05_scatter_isoprice.png", "Shrinking cut prices to S$1.45M; after 2019 they jump to S$2.05M"),
    ("06_violin_pricing.png", "Every unit size sells at its project's PSF, within 1.5%"),
    ("07_lollipop_macro.png", "Least affordable in 2007–13, when that land was sold"),
    ("08_heatmap_tdsr.png", "The median household can borrow for under 800 sq ft"),
    ("09_waterfall.png", "Price tag S$1.46M → S$2.18M, driven by PSF (+39%)"),
]
LABEL_H, HEAD_H, ROW_GAP = 0.18, 0.34, 0.05
PANEL_H = 0.05 + HEAD_H + 0.02 + FIG_H + INSIGHT_H + 0.06
# what each chart means, beyond what can be read straight off it
INSIGHT = [
    "New launches crowd into 500–900 sq ft. A buyer who wants a family-sized home now mostly has to look at resale.",
    "A 43% drop in five years is too fast for slowly changing tastes. It lines up with the least affordable years (chart 7) and stops at URA's line.",
    "The cap cut shoeboxes, yet big homes never returned. Developers moved to compact one- and two-bedders instead.",
    "The raw and controlled gaps are almost identical, so units shrank everywhere, in cheap suburbs and prime districts alike.",
    "Smaller homes stayed cheaper despite a higher PSF, the trade shrinkflation relies on. Once the shrink stopped, the price tag rose by S$0.6M.",
    "Small units carry no premium inside a launch, so every sq ft cut comes straight off the price tag. That made shrinking the easiest lever.",
    "When prices ran ahead of incomes, developers cut the size of the product instead of its price per sq ft.",
    "The median new unit sits almost exactly at what the median household can borrow, which suggests launches are sized to the loan limit.",
    "With the cap fixing sizes, any further PSF rise reaches buyers in full. Affordability now hinges on land costs and incomes.",
]
y = KY + KH + 0.07
for r, label in enumerate(ROWS):
    text(M, y, CW, LABEL_H, [([(f"{r + 1}   {label}", {"b": True})], False)], f"Row {r + 1} label", size=7, color=BLUE2, anchor=MSO_ANCHOR.MIDDLE)
    py = y + LABEL_H + 0.03
    for c in range(3):
        k = r * 3 + c
        fig, head = PANELS[k]
        px = M + c * (COL + GUT)
        box(px, py, COL, PANEL_H, CARD, f"Panel {k + 1} card")
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(px + 0.08), Inches(py + 0.1), Inches(0.22), Inches(0.22))
        badge.name = f"Panel {k + 1} number"
        badge.fill.solid(); badge.fill.fore_color.rgb = RGBColor.from_string(NAVY); badge.line.fill.background(); badge.shadow.inherit = False
        tf = badge.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        run = p.add_run(); run.text = str(k + 1); run.font.name = FONT; run.font.size = Pt(7.5); run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        text(px + 0.37, py + 0.05, COL - 0.45, HEAD_H, [([(head, {"b": True})], False)], f"Panel {k + 1} takeaway", size=7.8,
             anchor=MSO_ANCHOR.MIDDLE)
        pic = slide.shapes.add_picture(f"{FIG_DIR}/{fig}", Inches(px + (COL - FIG_W) / 2), Inches(py + 0.05 + HEAD_H + 0.02), Inches(FIG_W), Inches(FIG_H))
        pic.name = f"Panel {k + 1} figure"
        text(px + 0.12, py + 0.05 + HEAD_H + 0.02 + FIG_H + 0.03, COL - 0.24, INSIGHT_H, [([(INSIGHT[k], {})], False)],
             f"Panel {k + 1} insight", size=6.9, color=INK2)
    y = py + PANEL_H + ROW_GAP

# ---------- answer: concludes on the question ----------
CY, CH = y + 0.03, 0.98
box(M, CY, CW, CH, NAVY, "Answer card")
text(M + 0.16, CY + 0.08, CW - 0.32, 0.34, [([("Answer: ", {"b": True, "color": "f7b89c"}),
                                                ("yes, until URA stopped it. Homes shrank about 40% to hold the price tag down. "
                                                 "Now that shrinking is capped, the price tag rises instead.", {"b": True, "color": WHITE})], False)],
     "Answer line", size=9.2, anchor=MSO_ANCHOR.MIDDLE)
FINDINGS = [
    ("1", "Homes on land sold after 2008 are about 40% smaller, and big homes all but disappeared"),
    ("2", "It is not location. It works because a launch is priced per sq ft, so fewer sq ft means a lower price tag"),
    ("3", "Prices outran incomes when that land was sold. Since the 2019 cap the new-launch price tag is up 49%"),
]
fcw = (CW - 0.32 - 2 * 0.2) / 3
for k, (num, finding) in enumerate(FINDINGS):
    fx = M + 0.16 + k * (fcw + 0.2)
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(fx), Inches(CY + 0.53), Inches(0.2), Inches(0.2))
    dot.name = f"Finding {num} number"
    dot.fill.solid(); dot.fill.fore_color.rgb = RGBColor.from_string(ORANGE); dot.line.fill.background(); dot.shadow.inherit = False
    tf = dot.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run = p.add_run(); run.text = num; run.font.name = FONT; run.font.size = Pt(7); run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(WHITE)
    text(fx + 0.28, CY + 0.5, fcw - 0.28, 0.44, [([(finding, {"color": WHITE})], False)], f"Finding {num}", size=7.2)

# ---------- footnote: data, limitations, sources ----------
FY = CY + CH + 0.06
text(M, FY, CW, SH - FY - 0.12, [
    ([("Data and method: ", {"b": True}), ("24,131 URA private home sales in 2025 and 74,436 from Aug 2020 to Dec 2023 [1]; 155 rows outside 200–5,000 sq ft "
                                          "removed. Charts 2–5 use 99-year leaseholds grouped by lease-start year. Chart 7 uses the URA price index [2] and "
                                          "SingStat household income [3]; chart 8 applies MAS loan rules [4]. ", {}),
      ("Limitations: ", {"b": True}), ("eras are compared at 2025 prices, so older homes also carry age and lease decay; no bedroom count; 2024 missing; "
                                      "associations, not causal effects.", {})], False),
    ([("Sources: ", {"b": True}), ("[1] URA REALIS via DAO2702 (2025; Lecture 6, 2020–23). [2] URA Private Residential Price Index, data.gov.sg. "
                                  "[3] SingStat household income, data.gov.sg. [4] MAS: TDSR 55%, 4% stress rate. [5] URA circular DC18-06 (2018). "
                                  "[6] ONS (2019), shrinkflation and CPIH. Code: shrinkflation.ipynb.", {})], False),
], "Footnote", size=5, color=MUTED, space_after=1)

out = "shrinkflation_poster.pptx"
prs.save(out)
print("saved", out)
